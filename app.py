"""Image observations, human review and DOCX export using Snowflake Cortex."""
import io
import json
import os
import tempfile
import uuid
from pathlib import Path
import streamlit as st
from docx import Document
from docx.shared import Inches
from utils.snowflake_utils import get_session, query_cortex
from utils.prompt_utils import build_defect_code_prompt

st.set_page_config(page_title='Visual Quality Assistant', layout='wide')
st.title('Visual Quality Assistant')
st.caption('Image observations with editable notes. Demonstration categories; human review required.')
files = st.file_uploader('Images', type=['jpg', 'jpeg', 'png'], accept_multiple_files=True)
question = st.text_input('What should the assistant look for?', 'Describe visible marks, alignment and obstructions.')
st.info('Analysis uploads images to your configured Snowflake stage and calls Cortex. Configure retention on that stage. Use synthetic or authorised images.')
if st.button('Analyse images', disabled=not files):
    required = ['SNOWFLAKE_ACCOUNT', 'SNOWFLAKE_USER', 'SNOWFLAKE_PRIVATE_KEY_PATH', 'STAGE_NAME', 'MODEL_NAME']
    missing = [key for key in required if not os.getenv(key)]
    if missing:
        st.error('Missing configuration: ' + ', '.join(missing))
    else:
        session = None
        try:
            session = get_session()
            stage = '@' + os.environ['STAGE_NAME'].lstrip('@')
            codes = json.loads(Path('codes/defect_codes.json').read_text())
            results = []
            with tempfile.TemporaryDirectory() as directory:
                for upload in files:
                    filename = uuid.uuid4().hex + Path(upload.name).suffix.lower()
                    path = Path(directory) / filename
                    path.write_bytes(upload.getvalue())
                    session.file.put(str(path), stage, auto_compress=False, overwrite=False)
                    result = query_cortex(session, build_defect_code_prompt(question, codes), stage, filename)
                    results.append({'name': upload.name, 'image': upload.getvalue(), 'text': result})
            st.session_state['results'] = results
        except Exception:
            st.error('Analysis failed. Check credentials, stage permissions, model access and file support in your Snowflake account.')
        finally:
            if session is not None:
                session.close()
results = st.session_state.get('results', [])
if results:
    doc = Document()
    doc.add_heading('Image observation report', 0)
    doc.add_paragraph('AI-assisted draft. Review all observations before use.')
    for i, item in enumerate(results):
        st.image(item['image'], caption=item['name'], width=400)
        item['text'] = st.text_area('Review observations: ' + item['name'], item['text'], key='review-' + str(i))
        doc.add_heading(item['name'], 1)
        doc.add_picture(io.BytesIO(item['image']), width=Inches(5))
        doc.add_paragraph(item['text'])
    output = io.BytesIO()
    doc.save(output)
    st.download_button('Download reviewed DOCX', output.getvalue(), 'image-observations.docx')
